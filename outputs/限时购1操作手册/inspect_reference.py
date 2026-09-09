from pathlib import Path
from docx import Document

reference = Path(r"C:\Users\Administrator.DESKTOP-734BMB4\Desktop\发票管理操作手册-供应商平台.docx")
document = Document(reference)

print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} sections={len(document.sections)}")
for index, paragraph in enumerate(document.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        print(f"P{index}\t{paragraph.style.name}\t{text}")

for table_index, table in enumerate(document.tables, start=1):
    print(f"TABLE {table_index}: {len(table.rows)}x{len(table.columns)}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
