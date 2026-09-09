from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


REFERENCE = Path(r"C:\Users\Administrator.DESKTOP-734BMB4\Desktop\发票管理操作手册-供应商平台.docx")
OUTPUT = Path(r"D:\Github\-ERP-\outputs\限时购1操作手册\限时购1操作手册-供应商平台.docx")


def set_run_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_style_font(style, name, size, color, bold=False, before=0, after=0, line=1.35):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(*color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def clear_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def shade_paragraph(paragraph, fill="FFF2CC"):
    props = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def add_heading(document, text, level):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 16, 2: 14, 3: 12}[level], bold=True, color=(31, 78, 121))
    return paragraph


def add_body(document, text, bold_prefix=None, note=False):
    paragraph = document.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, size=10.5, bold=True, color=(31, 78, 121) if bold_prefix.startswith("✅") else (0, 0, 0))
        run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(run, size=10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5)
    if note:
        shade_paragraph(paragraph)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_path(document, text):
    add_body(document, "✅ 操作路径：", bold_prefix="✅ 操作路径：")
    return add_body(document, text)


def create_decimal_numbering(document):
    numbering = document.part.numbering_part.element
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_num = OxmlElement("w:abstractNumId")
    abstract_num.set(qn("w:val"), "1")
    number.append(abstract_num)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    number.append(override)
    numbering.append(number)
    return num_id


def add_steps(document, steps):
    num_id = create_decimal_numbering(document)
    for step in steps:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.space_after = Pt(5)
        props = paragraph._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        numbering.extend([level, number])
        props.append(numbering)
        run = paragraph.add_run(step)
        set_run_font(run, size=10.5)


def add_note(document, text):
    return add_body(document, f"注意事项：{text}", bold_prefix="注意事项：", note=True)


def configure_styles(document):
    set_style_font(document.styles["Normal"], "微软雅黑", 10.5, (0, 0, 0), after=6, line=1.35)
    set_style_font(document.styles["Title"], "微软雅黑", 18, (31, 78, 121), True, after=18, line=1.1)
    set_style_font(document.styles["Heading 1"], "微软雅黑", 16, (31, 78, 121), True, before=16, after=8, line=1.2)
    set_style_font(document.styles["Heading 2"], "微软雅黑", 14, (47, 85, 151), True, before=12, after=6, line=1.2)
    set_style_font(document.styles["Heading 3"], "微软雅黑", 12, (0, 0, 0), True, before=10, after=5, line=1.2)


def add_footer_object_alt_text(document):
    for section in document.sections:
        for doc_prop in section.footer._element.iter(qn("wp:docPr")):
            doc_prop.set("title", "供应商平台操作手册页脚装饰")
            doc_prop.set("descr", "供应商平台操作手册页脚装饰")


def build():
    document = Document(REFERENCE)
    clear_body(document)
    configure_styles(document)
    add_footer_object_alt_text(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("【操作手册】限时购1-供应商平台")
    set_run_font(title_run, size=18, bold=True, color=(31, 78, 121))

    subtitle = document.add_paragraph(style="Normal")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("适用范围：供应商平台营销中心 · 限时购1")
    set_run_font(subtitle_run, size=10.5, color=(89, 89, 89))
    subtitle.paragraph_format.space_after = Pt(16)

    add_heading(document, "核心功能概述", 1)
    add_body(document, "限时购1用于在指定活动时间内，为已选商品及参与规格配置限时价、限购数量和活动库存。供应商可在活动列表中查询、查看和编辑活动，并在商品维度或规格维度维护参与状态与促销配置。")
    add_body(document, "当前版本的限时购1面向全部买家，不再提供活动买家范围、指定买家分组、买家分组筛选或买家ID筛选。")
    add_note(document, "本手册依据当前供应商平台前端原型编写。保存动作目前完成页面必填校验；真实后端持久化、库存扣减与回退、买家端成交价计算及活动提前结束规则需以正式服务端方案为准。")

    add_heading(document, "操作流程说明", 1)
    add_heading(document, "限时购1", 2)

    add_heading(document, "1.1 限时购1列表", 3)
    add_body(document, "限时购1列表展示活动ID、活动名称、活动商品数、开始时间、结束时间、状态及操作。活动状态包括未开始、进行中、已结束。")
    add_path(document, "供应商平台：【营销】→【限时购1】")
    add_body(document, "☞操作流程：", bold_prefix="☞操作流程：")
    add_steps(document, [
        "进入【限时购1】页面。",
        "可按活动状态、活动时间、活动名称、活动ID、商品ID或规格ID填写查询条件。",
        "点击【查询】查看符合条件的活动；点击【重置】恢复默认筛选状态。",
        "在列表操作列点击【查看】进入活动详情，或点击【编辑】进入活动配置页。",
    ])
    add_note(document, "限时购1列表不展示活动买家范围，也不支持按买家分组或买家ID查询。")

    add_heading(document, "1.2 新增限时购1活动", 3)
    add_path(document, "供应商平台：【营销】→【限时购1】→【新增限时购】")
    add_body(document, "☞操作流程：", bold_prefix="☞操作流程：")
    add_steps(document, [
        "填写活动名称、活动分类、开始时间和结束时间。活动名称最多可输入10个字符。",
        "点击【+ 选择商品】，在商品选择弹窗中勾选需要参加活动的商品。",
        "确认选择后，在商品详情区域维护商品限时价、总限购数量和总活动库存，或进入规格编辑页按规格分别配置。",
        "确认参与规格及活动商品信息后，点击【保存】完成前端校验。",
    ])
    add_note(document, "活动时间、活动分类等字段应按业务规则完整填写。当前原型的保存校验重点覆盖商品限时价、总限购数量和总活动库存。")

    add_heading(document, "1.3 商品与规格配置", 3)
    add_body(document, "商品可使用商品维度的统一配置，也可在规格编辑中按规格设置。若某一字段切换为规格维度生效，对应商品维度字段将显示为“按规格维度生效”。")
    add_path(document, "新增或编辑活动：【商品详情】→【编辑】")
    add_body(document, "☞操作流程：", bold_prefix="☞操作流程：")
    add_steps(document, [
        "在活动商品详情中点击对应商品的【编辑】进入规格配置。",
        "勾选需批量操作的规格，或对单个规格填写限时价、限购数量和活动库存。",
        "可将可用规格加入活动，或将参与规格撤出活动。撤出后，该规格的活动价格、限购和活动库存配置会被清空。",
        "保存规格配置后返回商品详情，检查参与规格数量和字段生效维度。",
    ])
    add_note(document, "每个活动商品至少保留一个参与规格；若尝试撤出全部参与规格，系统会提示“请至少保留一个规格参与活动”。")

    add_heading(document, "1.4 保存校验", 3)
    add_body(document, "点击【保存】时，系统会检查每个活动商品的必要促销配置。")
    add_body(document, "需重点检查的字段：商品限时价、总限购数量、总活动库存。任一字段为空时，系统会在对应商品处标记，并提示“商品限时价、总限购数量、总活动库存为空，请检查”或其中缺失字段的组合提示。")
    add_note(document, "总限购数量用于控制单个买家ID最多购买数量；填写0表示不做数量限制的业务含义需以正式交易规则为准。")

    add_heading(document, "1.5 编辑活动与单品终止", 3)
    add_path(document, "供应商平台：【营销】→【限时购1】→活动操作列【编辑】")
    add_body(document, "☞操作流程：", bold_prefix="☞操作流程：")
    add_steps(document, [
        "在列表中找到目标活动，点击【编辑】进入编辑页。",
        "开始时间在编辑态不可修改；可维护结束时间及现有活动商品的促销配置。",
        "如需停止某个活动商品，在商品操作列点击【单品终止】，并在确认弹窗中确认操作。",
        "终止后，该商品状态显示为已失效，规格详情仍可查看。",
    ])
    add_note(document, "当前原型支持单品终止的页面状态展示。活动级“提前结束”按钮未形成可持久化的正式交易操作，不应据此认定活动已在服务端结束。")

    add_heading(document, "1.6 查看活动详情", 3)
    add_path(document, "供应商平台：【营销】→【限时购1】→活动操作列【查看】")
    add_body(document, "☞操作流程：", bold_prefix="☞操作流程：")
    add_steps(document, [
        "在活动列表中点击【查看】进入详情页。",
        "查看活动名称、活动分类、开始时间、结束时间及商品详情。",
        "查看商品商城价、限时价、总限购数量、活动总库存、规格数量及商品活动状态。",
        "点击规格数量列的【查看】可查看该商品的规格明细。",
    ])
    add_note(document, "限时购1详情页不展示活动买家范围或买家分组信息。")

    add_heading(document, "1.7 商品活动状态说明", 3)
    add_body(document, "商品活动状态用于辅助判断商品是否仍享受活动配置：活动未开始时显示“未生效”；活动进行中时显示“生效中”；活动已结束或商品被手动终止时显示“已失效”。")
    add_note(document, "商品状态以活动状态与商品终止状态为准。实际买家端展示、下单拦截和库存占用需要服务端在交易链路中再次校验。")

    add_heading(document, "使用注意事项", 1)
    add_body(document, "1. 限时购1已移除活动买家范围与买家分组功能，活动不再按买家分组配置或展示。")
    add_body(document, "2. 配置限时价、限购数量和活动库存前，请确认商品及规格参与状态，避免将未参与规格误配为活动规格。")
    add_body(document, "3. 本手册仅说明当前供应商平台操作界面。订单支付、取消、退款后的活动库存处理，以及买家端价格与限购提示，应以正式接口和交易规则为准。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
