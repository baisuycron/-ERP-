from copy import copy
from pathlib import Path

import openpyxl


BASE = Path(r"D:\Thunderobot\GitHub\-ERP-\outputs\km_invoice_prd_tests")
SOURCE = BASE / "发票二期测试用例-AI版.xlsx"
TEMPLATE = Path(
    r"C:\Users\Thunderobot\Documents\WXWork\1688857839786170\Cache\File\2026-06\发票二期测试用例1-20260527别动.xlsx"
)
OUTPUT = BASE / "发票二期测试用例-AI版-单表格式.xlsx"

TARGET_HEADERS = ["目录", "标题", "优先级", "描述", "前置条件", "步骤", "预期结果"]


def text(value):
    return "" if value is None else str(value).strip()


def get_row_dict(ws, row_idx, headers):
    return {header: ws.cell(row_idx, col_idx).value for col_idx, header in enumerate(headers, start=1)}


def build_directory(row):
    role = text(row.get("端/角色"))
    module = text(row.get("模块"))
    feature = text(row.get("功能点"))
    left = role or module
    if role and module and module not in role:
        left = f"{role}/{module}"
    if feature:
        return f"{left}-{feature}" if left else feature
    return left


def build_description(row):
    parts = []
    feature = text(row.get("功能点"))
    case_type = text(row.get("用例类型"))
    data = text(row.get("测试数据"))
    source = text(row.get("来源/备注"))
    if feature:
        parts.append(f"功能点：{feature}")
    if case_type:
        parts.append(f"用例类型：{case_type}")
    if data:
        parts.append(f"测试数据：{data}")
    if source:
        parts.append(f"来源：{source}")
    return "\n".join(parts)


src_wb = openpyxl.load_workbook(SOURCE, data_only=False)
src_ws = src_wb["全部用例"]
src_headers = [text(src_ws.cell(1, col_idx).value) for col_idx in range(1, src_ws.max_column + 1)]

tmpl_wb = openpyxl.load_workbook(TEMPLATE)
ws = tmpl_wb.active
ws.title = "发票二期"

header_styles = [copy(ws.cell(1, col_idx)._style) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
header_fills = [copy(ws.cell(1, col_idx).fill) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
header_fonts = [copy(ws.cell(1, col_idx).font) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
header_alignments = [copy(ws.cell(1, col_idx).alignment) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
header_borders = [copy(ws.cell(1, col_idx).border) for col_idx in range(1, len(TARGET_HEADERS) + 1)]

row_style_source = 295 if ws.max_row >= 295 else 2 if ws.max_row >= 2 else 1
row_style = [copy(ws.cell(row_style_source, col_idx)._style) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
row_fill = [copy(ws.cell(row_style_source, col_idx).fill) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
row_font = [copy(ws.cell(row_style_source, col_idx).font) for col_idx in range(1, len(TARGET_HEADERS) + 1)]
row_alignment = [
    copy(ws.cell(row_style_source, col_idx).alignment) for col_idx in range(1, len(TARGET_HEADERS) + 1)
]
row_border = [copy(ws.cell(row_style_source, col_idx).border) for col_idx in range(1, len(TARGET_HEADERS) + 1)]

# Clear existing rows while keeping the worksheet object, column widths, and page setup.
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)

for col_idx, header in enumerate(TARGET_HEADERS, start=1):
    ws.cell(1, col_idx).value = header

for col_idx in range(1, len(TARGET_HEADERS) + 1):
    cell = ws.cell(1, col_idx)
    cell._style = header_styles[col_idx - 1]
    cell.fill = header_fills[col_idx - 1]
    cell.font = header_fonts[col_idx - 1]
    cell.alignment = header_alignments[col_idx - 1]
    cell.border = header_borders[col_idx - 1]

target_row = 2
for row_idx in range(2, src_ws.max_row + 1):
    row = get_row_dict(src_ws, row_idx, src_headers)
    if not text(row.get("用例标题")):
        continue
    values = [
        build_directory(row),
        text(row.get("用例标题")),
        text(row.get("优先级")),
        build_description(row),
        text(row.get("前置条件")),
        text(row.get("操作步骤")),
        text(row.get("预期结果")),
    ]
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(target_row, col_idx)
        cell.value = value
        cell._style = copy(row_style[col_idx - 1])
        cell.fill = copy(row_fill[col_idx - 1])
        cell.font = copy(row_font[col_idx - 1])
        cell.alignment = copy(row_alignment[col_idx - 1])
        cell.border = copy(row_border[col_idx - 1])
        cell.alignment = copy(cell.alignment)
        cell.alignment = openpyxl.styles.Alignment(
            horizontal=cell.alignment.horizontal,
            vertical="top",
            text_rotation=cell.alignment.text_rotation,
            wrap_text=True,
            shrink_to_fit=cell.alignment.shrink_to_fit,
            indent=cell.alignment.indent,
        )
    target_row += 1

last_row = target_row - 1
ws.auto_filter.ref = f"A1:G{last_row}"
ws.freeze_panes = "A2"
ws.sheet_view.topLeftCell = "A1"
ws.sheet_view.selection[0].sqref = "A2"
ws.sheet_view.selection[0].activeCell = "A2"

for row_idx in range(2, last_row + 1):
    ws.row_dimensions[row_idx].height = 64

# Preserve the template's visible column proportions.
widths = {"A": 30, "B": 40.125, "C": 10, "D": 4.625, "E": 36.75, "F": 33.5, "G": 55.75}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

tmpl_wb.save(OUTPUT)
print(OUTPUT)
print(f"rows={last_row - 1}")
