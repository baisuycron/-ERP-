import json
import os
import re
from collections import Counter, defaultdict

import openpyxl


BASE = r"D:\Thunderobot\GitHub\-ERP-\outputs\km_invoice_prd_tests"
FILES = {
    "AI版": os.path.join(BASE, "ai_version_tmp.xlsx"),
    "20260527别动": os.path.join(BASE, "human_20260527_tmp.xlsx"),
}

CASE_HEADERS = [
    "用例标题",
    "用例名称",
    "标题",
    "测试点",
    "功能点",
]

KEYWORDS = {
    "任务中心": ["任务中心", "下载任务", "任务列表", "任务状态", "刷新", "关闭弹窗", "下载完成", "生成中", "排队", "等待"],
    "导出记录旧入口": ["导出记录", "导出纪录", "开票记录", "导出列表"],
    "批量下载/200限制": ["批量下载", "批量导出", "200", "二百", "勾选", "全选", "跨页"],
    "导入校验": ["导入", "上传", "xlsx", "zip", "pdf", "格式", "失败数据", "失败原因", "模板"],
    "小程序/跨页选择": ["小程序", "跨页", "分页", "全选", "当前页", "取消选择"],
    "边界规则": ["2026-01-01", "2026/01/01", "2026年1月1日", "售后", "货到付款", "COD", "0元", "零元", "金额为0"],
    "异步状态": ["生成中", "等待中", "排队", "处理中", "失败", "成功", "超时", "重新下载", "刷新"],
    "权限/角色": ["权限", "角色", "管理员", "普通", "无权限", "登录", "账号"],
    "待办": ["待办", "我的待办", "平台待办", "卖家待办"],
    "温馨提示": ["温馨提示", "提示语", "提示"],
}


def cell_text(value):
    return "" if value is None else str(value).strip()


def normalize(text):
    text = cell_text(text).lower()
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[\[\]【】()（）,，.。:：;；/\\|\-_—~～\"“”'‘’<>＞]+", "", text)
    for token in ["发票二期", "测试用例", "验证"]:
        text = text.replace(token, "")
    return text


def find_header_row(ws):
    best_score = -1
    best_row = 1
    for row_idx in range(1, min(ws.max_row or 1, 15) + 1):
        values = [cell_text(c.value) for c in ws[row_idx]]
        nonempty = [v for v in values if v]
        hits = sum(
            any(k in v for k in ["用例", "标题", "步骤", "预期", "优先级", "模块", "前置", "测试点", "操作"])
            for v in nonempty
        )
        score = hits * 10 + len(nonempty)
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def classify_cols(headers):
    cols = {}
    title_candidates = []
    for index, header in enumerate(headers, start=1):
        if not header:
            continue
        if any(k in header for k in CASE_HEADERS) and "预期" not in header:
            if "用例标题" in header or "用例名称" in header:
                rank = 0
            elif header == "标题":
                rank = 1
            elif "测试点" in header:
                rank = 2
            else:
                rank = 3
            title_candidates.append((rank, index))
        if any(k in header for k in ["测试步骤", "操作步骤", "步骤", "操作"]):
            cols.setdefault("steps", index)
        if any(k in header for k in ["预期结果", "期望结果", "预期"]):
            cols.setdefault("expected", index)
        if any(k in header for k in ["前置条件", "前提"]):
            cols.setdefault("precondition", index)
        if any(k in header for k in ["优先级", "等级"]):
            cols.setdefault("priority", index)
        if any(k in header for k in ["模块", "功能模块", "所属模块", "目录"]):
            cols.setdefault("module", index)
        if any(k in header for k in ["编号", "ID", "id", "序号", "用例ID"]):
            cols.setdefault("id", index)
    if title_candidates:
        cols["title"] = sorted(title_candidates)[0][1]
    return cols


def row_text(ws, row_idx, max_col):
    return "\n".join(cell_text(ws.cell(row_idx, col_idx).value) for col_idx in range(1, max_col + 1))


def read_workbook(label, path):
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    canonical_sheet = "全部用例" if "全部用例" in wb.sheetnames else None
    book = {
        "path": path,
        "sheets": [],
        "cases": [],
        "kw": {bucket: 0 for bucket in KEYWORDS},
        "kw_examples": defaultdict(list),
        "priorities": Counter(),
        "modules": Counter(),
    }

    for ws in wb.worksheets:
        include_as_case = canonical_sheet is None or ws.title == canonical_sheet
        max_row = ws.max_row or 0
        max_col = min(ws.max_column or 0, 40)
        if max_row == 0 or max_col == 0:
            continue
        header_row = find_header_row(ws)
        headers = [cell_text(ws.cell(header_row, col_idx).value) for col_idx in range(1, max_col + 1)]
        cols = classify_cols(headers)
        data_rows = 0
        sheet_cases = 0

        for row_idx in range(header_row + 1, max_row + 1):
            values = [cell_text(ws.cell(row_idx, col_idx).value) for col_idx in range(1, max_col + 1)]
            if not any(values):
                continue
            data_rows += 1
            text = "\n".join(values)

            # Skip summary rows and duplicate module sheets when a canonical case sheet exists.
            if not include_as_case:
                continue
            if "title" not in cols and not any(k in text for k in ["步骤", "预期", "PC商城", "小程序", "供应商", "平台中心"]):
                continue

            title = cell_text(ws.cell(row_idx, cols["title"]).value) if "title" in cols else ""
            if not title:
                nonnums = [v for v in values if v and not re.fullmatch(r"\d+(\.\d+)?", v)]
                title = nonnums[0] if nonnums else values[0]
            if not title:
                continue

            module = cell_text(ws.cell(row_idx, cols["module"]).value) if "module" in cols else ws.title
            priority = cell_text(ws.cell(row_idx, cols["priority"]).value) if "priority" in cols else ""
            case = {
                "sheet": ws.title,
                "row": row_idx,
                "title": title,
                "norm": normalize(title),
                "module": module,
                "priority": priority,
                "text": text,
            }
            book["cases"].append(case)
            sheet_cases += 1
            if priority:
                book["priorities"][priority] += 1
            if module:
                book["modules"][module] += 1

            lowered = text.lower()
            for bucket, terms in KEYWORDS.items():
                if any(term.lower() in lowered for term in terms):
                    book["kw"][bucket] += 1
                    if len(book["kw_examples"][bucket]) < 5:
                        book["kw_examples"][bucket].append({"sheet": ws.title, "row": row_idx, "title": title})

        book["sheets"].append(
            {
                "name": ws.title,
                "max_row": max_row,
                "max_col": ws.max_column or 0,
                "header_row": header_row,
                "headers": [h for h in headers if h],
                "cols": cols,
                "data_rows": data_rows,
                "case_rows": sheet_cases,
            }
        )

    title_counts = Counter(c["norm"] for c in book["cases"] if c["norm"])
    book["case_count"] = len(book["cases"])
    book["unique_norm_titles"] = len(title_counts)
    book["duplicate_norm_titles"] = sum(1 for _, count in title_counts.items() if count > 1)
    return book


def examples(book, norms, limit=25):
    output = []
    seen = set()
    for case in book["cases"]:
        if case["norm"] in norms and case["norm"] not in seen:
            output.append(
                {
                    "sheet": case["sheet"],
                    "row": case["row"],
                    "module": case["module"],
                    "priority": case["priority"],
                    "title": case["title"],
                }
            )
            seen.add(case["norm"])
            if len(output) >= limit:
                break
    return output


result = {label: read_workbook(label, path) for label, path in FILES.items()}
sets = {label: {case["norm"] for case in book["cases"] if case["norm"]} for label, book in result.items()}
common = sets["AI版"] & sets["20260527别动"]
only_ai = sets["AI版"] - sets["20260527别动"]
only_human = sets["20260527别动"] - sets["AI版"]

summary = {
    "books": {
        label: {
            key: value
            for key, value in book.items()
            if key not in ["cases", "kw_examples", "priorities", "modules"]
        }
        for label, book in result.items()
    },
    "priority_counts": {label: dict(book["priorities"]) for label, book in result.items()},
    "module_counts_top": {label: book["modules"].most_common(30) for label, book in result.items()},
    "keyword_examples": {label: dict(book["kw_examples"]) for label, book in result.items()},
    "exact_norm_common": len(common),
    "only_ai_count": len(only_ai),
    "only_20260527_count": len(only_human),
    "only_ai_examples": examples(result["AI版"], only_ai),
    "only_20260527_examples": examples(result["20260527别动"], only_human),
}

print(json.dumps(summary, ensure_ascii=False, indent=2))
