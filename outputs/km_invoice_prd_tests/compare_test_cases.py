import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

BASE = Path("D:/Thunderobot/GitHub/-ERP-/outputs/km_invoice_prd_tests")
FILES = {
    "Codex-AI": BASE / "codex_ai.xlsx",
    "GLM": BASE / "glm.xlsx",
    "OtherAI": BASE / "other_ai.xlsx",
}

HEADER_ALIASES = {
    "id": ["用例ID", "编号", "序号", "ID"],
    "module": ["模块", "所属模块", "功能模块", "测试模块", "一级模块"],
    "client": ["端", "端/角色", "角色", "适用端", "测试端"],
    "feature": ["功能点", "功能", "子功能", "测试点", "测试项"],
    "title": ["用例标题", "测试用例", "用例名称", "测试标题", "场景", "标题"],
    "precondition": ["前置条件", "前置", "预置条件"],
    "steps": ["操作步骤", "测试步骤", "步骤"],
    "data": ["测试数据", "数据"],
    "expected": ["预期结果", "期望结果", "预期"],
    "priority": ["优先级", "P级"],
    "type": ["用例类型", "类型", "测试类型"],
}

KEYWORDS = {
    "PC发票管理": ["PC", "商家中心", "已开具", "可申请", "已申请"],
    "任务中心/导出下载": ["任务中心", "导出", "下载", "异步", "执行成功", "执行失败"],
    "批量下载发票": ["批量下载", "压缩包", "发票PDF", "下载状态", "已下载", "未下载"],
    "小程序发票管理": ["小程序", "发票管理", "抬头管理", "可申请开票", "已申请开票", "已开具发票"],
    "小程序批量申请": ["批量申请", "本页全选", "跨页", "异常订单", "提交申请"],
    "撤销申请": ["撤销", "查看发票", "二次确认", "已撤销", "重新提交"],
    "平台待办": ["平台", "待办管理", "即将超时", "已超时", "催办", "删除"],
    "卖家待办": ["卖家", "我的待办", "首页"],
    "批量导入发票": ["批量导入", "导入任务", "ZIP", "XLSX", "PDF", "失败数据"],
    "开票规则边界": ["售后", "先货后款", "订单总额", "2026-01-01", "支付时间", "未还款"],
    "权限隔离": ["权限", "账号", "仅当前", "隔离", "不可见"],
    "并发/状态反写": ["并发", "服务端校验", "状态变更", "反写", "失败不更新"],
}


def cell_text(v):
    if v is None:
        return ""
    return str(v).strip()


def norm(s):
    return re.sub(r"\s+", "", cell_text(s)).lower()


def find_header(rows):
    best = None
    best_score = -1
    for idx, row in enumerate(rows[:20]):
        values = [cell_text(c) for c in row]
        score = 0
        for aliases in HEADER_ALIASES.values():
            if any(any(a in v for a in aliases) for v in values):
                score += 1
        if score > best_score:
            best_score = score
            best = (idx, values)
    if best_score < 2:
        return None, []
    return best


def column_map(header):
    mapping = {}
    for key, aliases in HEADER_ALIASES.items():
        for i, name in enumerate(header):
            clean = cell_text(name)
            if any(alias == clean or alias in clean for alias in aliases):
                mapping[key] = i
                break
    return mapping


def parse_cases(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    cases = []
    preferred_sheet = "全部用例" if "全部用例" in wb.sheetnames else None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        nonempty_rows = [r for r in rows if any(cell_text(c) for c in r)]
        header_idx, header = find_header(nonempty_rows)
        sheets.append({
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "nonempty_rows": len(nonempty_rows),
            "header": header[:14],
        })
        if preferred_sheet and ws.title != preferred_sheet:
            continue
        if header_idx is None:
            continue
        cmap = column_map(header)
        # Ignore summary sheets if no title/steps/expected-like columns.
        if not ({"title", "steps", "expected"} & set(cmap)):
            continue
        for r in nonempty_rows[header_idx + 1:]:
            row_text = [cell_text(c) for c in r]
            if not any(row_text):
                continue
            item = {k: row_text[i] if i < len(row_text) else "" for k, i in cmap.items()}
            combined = " ".join(row_text)
            # Skip matrix/flow rows that are not executable test cases.
            has_case_signal = any(item.get(k) for k in ["title", "steps", "expected"])
            if not has_case_signal:
                continue
            item["_sheet"] = ws.title
            item["_combined"] = combined
            cases.append(item)
    return sheets, cases


def classify(text):
    hits = []
    for category, keys in KEYWORDS.items():
        if any(k.lower() in text.lower() for k in keys):
            hits.append(category)
    return hits or ["未归类"]


def richness(case):
    steps = case.get("steps", "")
    expected = case.get("expected", "")
    pre = case.get("precondition", "")
    data = case.get("data", "")
    return {
        "steps_len": len(steps),
        "expected_len": len(expected),
        "pre_len": len(pre),
        "data_len": len(data),
        "has_priority": bool(case.get("priority")),
        "has_type": bool(case.get("type")),
        "step_count": len(re.findall(r"(^|\n)\s*\d+[.、]", steps)),
        "expected_count": len(re.findall(r"(^|\n)\s*[-•]|(^|\n)\s*\d+[.、]", expected)),
    }


result = {}
for name, path in FILES.items():
    sheets, cases = parse_cases(path)
    cat_counter = Counter()
    module_counter = Counter()
    priority_counter = Counter()
    type_counter = Counter()
    rich = []
    titles = []
    for case in cases:
        text = case.get("_combined", "")
        for cat in classify(text):
            cat_counter[cat] += 1
        module_counter[case.get("module") or case.get("_sheet") or ""] += 1
        priority_counter[case.get("priority") or ""] += 1
        type_counter[case.get("type") or ""] += 1
        rich.append(richness(case))
        titles.append(case.get("title") or case.get("feature") or text[:80])
    avg = {}
    if rich:
        for key in rich[0].keys():
            vals = [r[key] for r in rich]
            avg[key] = round(sum(vals) / len(vals), 2)
    result[name] = {
        "path": str(path),
        "sheets": sheets,
        "case_count": len(cases),
        "categories": dict(cat_counter),
        "modules": dict(module_counter),
        "priorities": dict(priority_counter),
        "types": dict(type_counter),
        "avg_richness": avg,
        "sample_titles": titles[:12],
    }

all_categories = sorted(KEYWORDS.keys())
matrix = []
for cat in all_categories:
    row = {"category": cat}
    for name in FILES:
        row[name] = result[name]["categories"].get(cat, 0)
    matrix.append(row)
result["coverage_matrix"] = matrix

out = BASE / "test_case_comparison_summary.json"
out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
