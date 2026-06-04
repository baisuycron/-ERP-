from pathlib import Path

import openpyxl

BASE = Path("D:/Thunderobot/GitHub/-ERP-/outputs/km_invoice_prd_tests")
FILES = {
    "Codex-AI": BASE / "codex_ai.xlsx",
    "GLM": BASE / "glm.xlsx",
    "OtherAI": BASE / "other_ai.xlsx",
}
KEYWORDS = [
    "导出记录",
    "任务中心",
    "批量导入",
    "撤销",
    "下载状态",
    "本页全选",
    "即将超时",
    "已超时",
    "先货后款",
    "2026-01-01",
    "200",
    "失败数据",
    "服务端校验",
    "ZIP",
    "XLSX",
    "PDF",
]


def text(v):
    return "" if v is None else str(v)


for name, path in FILES.items():
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["全部用例"] if "全部用例" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [text(c).strip() for c in rows[0]]
    print(f"\n## {name} / {ws.title}")
    print("header:", header)
    for kw in KEYWORDS:
        matches = []
        for r in rows[1:]:
            combined = " ".join(text(c) for c in r)
            if kw in combined:
                title = text(r[4]) if len(r) > 4 and header[4] == "用例标题" else (text(r[1]) if len(r) > 1 else combined)
                matches.append(title[:100])
        print(f"{kw}: {len(matches)}" + ((" | " + "；".join(matches[:6])) if matches else ""))
