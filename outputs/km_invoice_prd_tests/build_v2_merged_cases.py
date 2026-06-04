import json
import re
from pathlib import Path

import openpyxl

BASE = Path("D:/Thunderobot/GitHub/-ERP-/outputs/km_invoice_prd_tests")
CODEX = BASE / "codex_ai.xlsx"
GLM = BASE / "glm.xlsx"
OTHER = BASE / "other_ai.xlsx"
OUT = BASE / "merged_v2_cases.json"


def text(v):
    return "" if v is None else str(v).strip()


def steps(*items):
    return "\n".join(f"{i + 1}. {item}" for i, item in enumerate(items))


def expected(*items):
    return "\n".join(f"- {item}" for item in items)


def normalize_title(s):
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"验证|功能正常|展示正确|正确|成功|失败", "", s)
    return s.lower()


def infer_module(text_blob):
    rules = [
        ("批量导入发票", ["批量导入", "导入任务", "ZIP", "XLSX", "PDF", "失败数据"]),
        ("任务中心", ["任务中心", "导出任务", "下载任务", "异步"]),
        ("订单撤销开票", ["订单>查看发票", "查看发票", "撤销申请", "撤销"]),
        ("平台待办", ["平台侧", "待办管理", "催办", "删除", "即将超时", "已超时"]),
        ("卖家待办", ["卖家侧", "我的待办", "首页"]),
        ("小程序发票管理", ["小程序", "本页全选", "批量申请", "已申请开票", "可申请开票"]),
        ("PC发票管理", ["PC商城", "PC", "商家中心", "发票下载状态", "批量下载发票"]),
        ("规则与边界", ["先货后款", "2026-01-01", "订单总额为0", "售后"]),
    ]
    for module, keys in rules:
        if any(k in text_blob for k in keys):
            return module
    return "补充回归"


def canonical_module(module):
    mapping = {
        "小程序订单查看发票": "订单撤销开票",
        "平台侧待办管理": "平台待办",
        "卖家侧我的待办": "卖家待办",
        "开票规则与提示": "规则与边界",
    }
    return mapping.get(module, module)


def infer_feature(text_blob, module):
    rules = [
        ("任务中心入口/列表", ["任务中心", "导出任务", "下载任务"]),
        ("批量下载发票", ["批量下载发票", "200", "压缩包"]),
        ("发票下载状态", ["下载状态", "已下载", "未下载"]),
        ("批量申请开票", ["批量申请", "提交申请", "异常订单", "编辑"]),
        ("本页全选/跨页选择", ["本页全选", "跨页", "已选"]),
        ("撤销申请", ["撤销"]),
        ("待办生成/处理", ["待办", "超时", "催办", "删除"]),
        ("批量导入校验", ["ZIP", "XLSX", "PDF", "失败数据", "导入"]),
        ("开票资格规则", ["售后", "先货后款", "2026-01-01", "订单总额"]),
    ]
    for feature, keys in rules:
        if any(k in text_blob for k in keys):
            return feature
    return module


def infer_client(module, text_blob):
    if module == "PC发票管理" or module == "任务中心":
        return "买家侧PC商城"
    if module in ["小程序发票管理", "订单撤销开票"]:
        return "买家侧小程序"
    if module == "平台待办":
        return "平台侧"
    if module in ["卖家待办", "批量导入发票"]:
        return "卖家侧/任务中心" if module == "批量导入发票" else "卖家侧"
    if "PC" in text_blob:
        return "买家侧PC商城"
    if "小程序" in text_blob:
        return "买家侧小程序"
    return "多端"


def infer_type(text_blob):
    if any(k in text_blob for k in ["超过", "边界", "200", "2026-01-01", "1M", "200M"]):
        return "边界"
    if any(k in text_blob for k in ["失败", "不可", "未选", "无数据", "错误", "缺失", "校验"]):
        return "异常"
    if any(k in text_blob for k in ["权限", "账号", "不可见"]):
        return "权限"
    if any(k in text_blob for k in ["并发", "服务端再校验", "状态变化"]):
        return "并发/异常"
    if any(k in text_blob for k in ["展示", "文案", "按钮", "入口"]):
        return "UI"
    return "功能"


def read_codex():
    wb = openpyxl.load_workbook(CODEX, read_only=True, data_only=True)
    ws = wb["全部用例"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [text(c) for c in rows[0]]
    out = []
    for row in rows[1:]:
        if not any(text(c) for c in row):
            continue
        item = dict(zip(headers, [text(c) for c in row]))
        module = canonical_module(item.get("模块", ""))
        out.append({
            "module": module,
            "client": infer_client(module, " ".join([item.get("模块", ""), item.get("端", "")])),
            "feature": item.get("功能点", ""),
            "title": item.get("用例标题", ""),
            "precondition": item.get("前置条件", ""),
            "steps": item.get("操作步骤", ""),
            "data": item.get("测试数据", ""),
            "expected": item.get("预期结果", ""),
            "priority": item.get("优先级", "P1"),
            "type": item.get("用例类型", "功能"),
            "source": "Codex-AI V1",
        })
    return out


def read_flat_ai(path, source):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for row in rows[1:]:
        if not any(text(c) for c in row):
            continue
        directory, title, priority, pre, step, exp = [text(c) for c in row[:6]]
        blob = " ".join([directory, title, pre, step, exp])
        # Old PRD path: page-level export-record entry was removed by 2026-05-26 update.
        if "导出记录" in blob and "任务中心" not in blob:
            if any(k in title for k in ["存在导出记录入口", "弹出导出记录弹窗", "导出记录弹窗", "点击导出记录"]):
                continue
            title = title.replace("导出记录", "任务中心导出/下载任务")
            step = step.replace("导出记录", "任务中心")
            exp = exp.replace("导出记录", "任务中心")
        module = canonical_module(infer_module(blob))
        feature = infer_feature(blob, module)
        out.append({
            "module": module,
            "client": infer_client(module, blob),
            "feature": feature,
            "title": title,
            "precondition": pre,
            "steps": step,
            "data": "",
            "expected": exp,
            "priority": priority or "P1",
            "type": infer_type(blob),
            "source": source,
        })
    return out


manual = [
    {
        "module": "任务中心",
        "client": "买家侧PC商城",
        "feature": "任务中心入口/列表",
        "title": "发票管理页点击导出后弹出任务中心引导提示",
        "precondition": "买家已进入PC商城商家中心发票管理页，当前页有可导出数据。",
        "steps": steps("在可申请/已申请/已开具任一页签点击导出。", "观察页面提示。", "点击提示中的查看任务中心或手动打开任务中心。"),
        "data": "任一发票页签查询结果。",
        "expected": expected("点击导出后立即生成异步任务。", "页面提示用户可在任务中心查看导出进度和下载文件。", "任务中心能看到刚生成的任务。"),
        "priority": "P0",
        "type": "功能",
        "source": "V2人工补充-最新PRD",
    },
    {
        "module": "任务中心",
        "client": "买家侧PC商城",
        "feature": "任务中心入口/列表",
        "title": "任务中心关闭按钮仅关闭弹窗不取消任务",
        "precondition": "任务中心存在处理中导出任务。",
        "steps": steps("打开任务中心。", "点击取消按钮关闭。", "再次打开任务中心。", "等待任务完成。"),
        "data": "处理中导出任务1条。",
        "expected": expected("取消按钮仅关闭任务中心页面。", "任务继续执行，不被取消。", "再次打开可查看最新任务状态。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2人工补充-图文规则",
    },
    {
        "module": "任务中心",
        "client": "买家侧PC商城",
        "feature": "任务中心入口/列表",
        "title": "任务中心X按钮仅关闭弹窗不取消任务",
        "precondition": "任务中心存在处理中下载任务。",
        "steps": steps("打开任务中心。", "点击右上角X关闭。", "再次打开任务中心。"),
        "data": "处理中发票文件下载任务。",
        "expected": expected("X按钮仅关闭弹窗。", "任务仍存在且状态持续更新。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2人工补充-图文规则",
    },
    {
        "module": "任务中心",
        "client": "买家侧PC商城",
        "feature": "任务中心入口/列表",
        "title": "任务中心无任务空态展示",
        "precondition": "新买家账号未发起过任何导出/下载/导入任务。",
        "steps": steps("登录买家账号并进入商家中心。", "点击任务中心。", "分别查看导出/下载任务TAB和导入任务TAB。"),
        "data": "新账号或清空任务数据账号。",
        "expected": expected("任务中心正常打开。", "两个TAB均展示空态，不报错。", "空态文案不误导用户执行下载。"),
        "priority": "P2",
        "type": "UI",
        "source": "V2人工补充-空态",
    },
    {
        "module": "PC发票管理",
        "client": "买家侧PC商城",
        "feature": "发票下载状态",
        "title": "发票下载状态筛选无数据时展示空态",
        "precondition": "当前账号不存在符合某一下载状态条件的已开具发票。",
        "steps": steps("进入已开具发票页签。", "选择发票下载状态=已下载或未下载。", "点击查询。"),
        "data": "选择一个当前无数据的下载状态。",
        "expected": expected("列表展示空态。", "分页、合计金额、选择框状态正确。", "不展示上一轮查询残留数据。"),
        "priority": "P2",
        "type": "UI",
        "source": "V2吸收-OtherAI细化",
    },
    {
        "module": "PC发票管理",
        "client": "买家侧PC商城",
        "feature": "批量下载发票",
        "title": "批量下载发票选中200笔订单成功生成任务",
        "precondition": "已开具发票列表存在至少200笔可下载订单。",
        "steps": steps("通过分页累计选择200笔订单。", "确认已选笔数为200。", "点击批量下载发票。", "打开任务中心查看任务。"),
        "data": "200笔已开具发票订单。",
        "expected": expected("允许选择200笔。", "成功生成发票文件下载任务。", "任务数据范围为200笔订单。"),
        "priority": "P0",
        "type": "边界",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "PC发票管理",
        "client": "买家侧PC商城",
        "feature": "批量下载发票",
        "title": "批量下载发票文件缺失时任务失败且状态不反写",
        "precondition": "选中订单中存在发票PDF缺失或文件服务异常。",
        "steps": steps("选中存在文件异常的已开具订单。", "点击批量下载发票。", "等待任务中心任务结束。", "回到已开具发票页查询下载状态。"),
        "data": "1笔PDF缺失订单。",
        "expected": expected("任务执行失败或部分失败。", "失败订单不更新为已下载。", "任务中心展示失败状态或失败原因。"),
        "priority": "P0",
        "type": "异常",
        "source": "V2人工补充-状态一致性",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "本页全选/跨页选择",
        "title": "滚动到新页时本页全选状态重新计算",
        "precondition": "小程序可申请开票列表超过1页。",
        "steps": steps("第一页点击本页全选。", "继续滚动加载第二页。", "观察第二页本页全选控件。"),
        "data": "第一页10笔可选，第二页10笔可选。",
        "expected": expected("第一页已选状态保留。", "第二页本页全选按第二页数据重新计算，默认未全选。", "底部已选笔数仍包含第一页已选订单。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "本页全选/跨页选择",
        "title": "未加载数据不参与本页全选",
        "precondition": "可申请开票列表总数超过当前已加载数据。",
        "steps": steps("进入可申请开票页，仅加载第一页。", "点击本页全选。", "查看已选笔数。"),
        "data": "总数30笔，当前加载10笔。",
        "expected": expected("仅当前已加载页的可选订单被选中。", "未加载的20笔不计入已选笔数。"),
        "priority": "P1",
        "type": "边界",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "筛选",
        "title": "小程序订单号精确搜索",
        "precondition": "存在可申请开票订单。",
        "steps": steps("进入可申请开票页筛选。", "输入完整订单号。", "点击查询。"),
        "data": "完整订单号：20260212022895768。",
        "expected": expected("仅返回匹配订单号的订单。", "订单号、金额、店铺信息与原订单一致。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2人工补充-字段筛选",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "筛选",
        "title": "小程序店铺名称模糊搜索",
        "precondition": "存在多个店铺订单。",
        "steps": steps("进入可申请开票页筛选。", "输入店铺名称关键字。", "点击查询。"),
        "data": "店铺名称关键字：松鼠。",
        "expected": expected("返回店铺名称包含关键字的订单。", "不匹配店铺不展示。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2人工补充-字段筛选",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "批量申请开票",
        "title": "批量申请开票编辑模式展示全选和移除按钮",
        "precondition": "已进入批量申请开票页。",
        "steps": steps("点击编辑按钮。", "观察页面控件。", "选择部分订单并点击移除。"),
        "data": "3笔订单。",
        "expected": expected("进入编辑模式后订单列表前展示复选框。", "页面展示本页全选和移除按钮。", "被移除订单不再参与本次提交。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "批量申请开票",
        "title": "批量申请开票备注超长校验",
        "precondition": "进入批量申请开票页。",
        "steps": steps("在批量备注输入超过限制的文本。", "点击提交申请。"),
        "data": "备注长度超过产品/后端限制。",
        "expected": expected("页面阻止提交或截断前明确提示。", "提示备注长度超限。", "不生成异常开票申请。"),
        "priority": "P1",
        "type": "边界",
        "source": "V2人工补充-字段校验",
    },
    {
        "module": "小程序发票管理",
        "client": "买家侧小程序",
        "feature": "批量申请开票",
        "title": "批量申请开票必填抬头缺失校验",
        "precondition": "进入批量申请开票页，至少1笔订单未选择发票抬头。",
        "steps": steps("不选择发票抬头。", "点击提交申请。"),
        "data": "缺少发票抬头。",
        "expected": expected("提交失败。", "异常订单中标记缺少发票抬头。", "用户补充抬头后可重新提交。"),
        "priority": "P0",
        "type": "异常",
        "source": "V2人工补充-字段校验",
    },
    {
        "module": "订单撤销开票",
        "client": "买家侧小程序",
        "feature": "撤销申请",
        "title": "订单详情撤销申请后可申请列表恢复该订单",
        "precondition": "订单已申请开票且卖家未开票。",
        "steps": steps("进入订单详情 > 查看发票。", "点击撤销申请并确认。", "进入发票管理 > 可申请开票。", "搜索该订单。"),
        "data": "待开票订单。",
        "expected": expected("撤销成功后订单状态为已撤销。", "该订单重新出现在可申请开票列表。", "可再次提交开票申请。"),
        "priority": "P0",
        "type": "功能",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "平台待办",
        "client": "平台侧",
        "feature": "待办生成/处理",
        "title": "无超时开票申请时0点不生成发票待办",
        "precondition": "系统内不存在即将超时或已超时开票申请。",
        "steps": steps("触发每天0点待办生成任务。", "进入平台侧店铺 > 待办管理。", "筛选发票管理。"),
        "data": "无超时/即将超时申请。",
        "expected": expected("不生成发票待办。", "待办列表为空或不新增发票待办。"),
        "priority": "P1",
        "type": "边界",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "平台待办",
        "client": "平台侧",
        "feature": "待办生成/处理",
        "title": "发票待办标题区分即将超时和已超时",
        "precondition": "同时存在即将超时和已超时开票申请。",
        "steps": steps("触发0点待办生成。", "进入平台侧待办管理。", "查看发票待办标题。"),
        "data": "即将超时申请1笔，已超时申请1笔。",
        "expected": expected("生成即将超时的开票申请提醒。", "生成已超时的开票申请提醒。", "标题文案和类型不混淆。"),
        "priority": "P0",
        "type": "功能",
        "source": "V2吸收-GLM细化",
    },
    {
        "module": "卖家待办",
        "client": "卖家侧",
        "feature": "待办生成/处理",
        "title": "卖家首页我的待办去处理跳转并筛选发票待办",
        "precondition": "卖家首页存在发票相关待办。",
        "steps": steps("登录卖家侧首页。", "点击我的待办中的发票待办去处理。", "观察跳转后的筛选条件和列表。"),
        "data": "已超时发票待办。",
        "expected": expected("跳转到对应发票待办/开票申请处理页面。", "列表默认筛选到对应待办类型。", "待办数量与首页一致。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2吸收-OtherAI细化",
    },
    {
        "module": "批量导入发票",
        "client": "卖家侧/任务中心",
        "feature": "批量导入校验",
        "title": "导入模板必填字段缺失生成失败数据",
        "precondition": "卖家进入批量导入发票页面。",
        "steps": steps("准备ZIP，模板中缺失发票号码或订单号必填字段。", "上传ZIP。", "等待导入任务完成。", "下载失败数据。"),
        "data": "模板缺失发票号码。",
        "expected": expected("任务状态为全部失败或部分失败。", "失败数据中标明必填字段缺失。", "订单开票状态不被错误更新。"),
        "priority": "P0",
        "type": "异常",
        "source": "V2人工补充-导入字段校验",
    },
    {
        "module": "批量导入发票",
        "client": "卖家侧/任务中心",
        "feature": "批量导入校验",
        "title": "导入模板发票号码重复校验",
        "precondition": "卖家进入批量导入发票页面。",
        "steps": steps("准备ZIP，模板中同一订单重复填写同一发票号码。", "上传ZIP。", "查看导入任务结果。"),
        "data": "重复发票号码：2024041810004212。",
        "expected": expected("重复数据按规则失败或幂等处理。", "不生成重复开票记录。", "失败时可在失败数据中定位重复行。"),
        "priority": "P1",
        "type": "边界",
        "source": "V2人工补充-导入字段校验",
    },
    {
        "module": "批量导入发票",
        "client": "卖家侧/任务中心",
        "feature": "批量导入校验",
        "title": "ZIP内包含无关文件时校验处理",
        "precondition": "卖家进入批量导入发票页面。",
        "steps": steps("准备ZIP，除模板和PDF文件夹外包含txt或图片文件。", "上传ZIP。", "查看校验结果。"),
        "data": "ZIP包含readme.txt。",
        "expected": expected("系统按规则忽略无关文件或提示文件结构不合法。", "处理结果清晰，不出现未知异常。"),
        "priority": "P2",
        "type": "边界",
        "source": "V2人工补充-导入文件结构",
    },
    {
        "module": "规则与边界",
        "client": "买家侧PC/小程序",
        "feature": "开票资格规则",
        "title": "支付时间为2025-12-31 23:59:59不可开票",
        "precondition": "订单已完成且无售后，但支付时间早于2026-01-01。",
        "steps": steps("进入可申请开票页。", "搜索该订单。", "尝试提交开票申请。"),
        "data": "支付时间=2025-12-31 23:59:59。",
        "expected": expected("订单不支持提交开票。", "页面提示支付时间在2026年1月1日前不支持开票。"),
        "priority": "P0",
        "type": "边界",
        "source": "V2人工补充-PRD温馨提示",
    },
    {
        "module": "规则与边界",
        "client": "买家侧PC/小程序",
        "feature": "开票资格规则",
        "title": "售后关闭订单可按无售后规则提交开票",
        "precondition": "订单已完成，售后状态为售后关闭，且满足其他开票条件。",
        "steps": steps("进入可申请开票页。", "搜索售后关闭订单。", "选择并提交开票申请。"),
        "data": "售后状态=售后关闭。",
        "expected": expected("订单可被选择。", "可提交开票申请。", "售后状态展示为售后关闭。"),
        "priority": "P1",
        "type": "功能",
        "source": "V2人工补充-售后枚举",
    },
]


cases = []
seen = set()


def add_case(case):
    key = (case["module"], normalize_title(case["title"]))
    if key in seen:
        return
    seen.add(key)
    cases.append(case)


for case in read_codex():
    add_case(case)

for path, source in [(GLM, "GLM吸收"), (OTHER, "OtherAI吸收")]:
    for case in read_flat_ai(path, source):
        # Avoid too-generic old cases and cases already superseded by V2 manual scenarios.
        blob = " ".join([case["module"], case["feature"], case["title"], case["steps"], case["expected"]])
        if "导出记录" in blob and "任务中心" not in blob:
            continue
        if source == "OtherAI吸收":
            unique_keywords = ["无数据", "去处理", "超过200", "状态筛选无数据", "跳转", "筛选正常"]
            if not any(k in blob for k in unique_keywords):
                continue
        add_case(case)

for case in manual:
    add_case(case)

prefixes = {
    "PC发票管理": "PC",
    "小程序发票管理": "MP",
    "订单撤销开票": "ORDER",
    "平台待办": "TODO",
    "卖家待办": "SELLER",
    "任务中心": "TASK",
    "批量导入发票": "IMPORT",
    "规则与边界": "RULE",
    "补充回归": "REG",
}
counters = {}
for case in cases:
    prefix = prefixes.get(case["module"], "CASE")
    counters[prefix] = counters.get(prefix, 0) + 1
    case["id"] = f"{prefix}-{counters[prefix]:03d}"
    case["status"] = "未执行"

summary = {
    "total": len(cases),
    "by_module": {},
    "by_source": {},
    "by_priority": {},
}
for case in cases:
    summary["by_module"][case["module"]] = summary["by_module"].get(case["module"], 0) + 1
    summary["by_source"][case["source"]] = summary["by_source"].get(case["source"], 0) + 1
    summary["by_priority"][case["priority"]] = summary["by_priority"].get(case["priority"], 0) + 1

OUT.write_text(json.dumps({"summary": summary, "cases": cases}, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=2))
